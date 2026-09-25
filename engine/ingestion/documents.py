from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


@dataclass(frozen=True)
class ExtractedDocument:
    path: Path
    kind: str
    text: str
    page_or_paragraph_count: int


def _clean_text(text: str) -> str:
    lines = [line.strip() for line in text.replace("\x00", "").splitlines()]
    cleaned: list[str] = []
    previous_blank = False
    for line in lines:
        is_blank = not line
        if is_blank and previous_blank:
            continue
        cleaned.append(line)
        previous_blank = is_blank
    return "\n".join(cleaned).strip()


def _extract_pdf(path: Path) -> ExtractedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("PDF desteği için 'pypdf' kurulmalı: pip install -r requirements.txt") from exc

    reader = PdfReader(str(path))
    parts: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        page_text = (page.extract_text() or "").strip()
        if page_text:
            parts.append(f"[Sayfa {index}]\n{page_text}")

    text = _clean_text("\n\n".join(parts))
    if not text:
        raise ValueError(
            "PDF'den metin çıkarılamadı. Dosya taranmış/görüntü tabanlı olabilir; OCR bu sürümde yok."
        )
    return ExtractedDocument(path, "pdf", text, len(reader.pages))


def _extract_docx(path: Path) -> ExtractedDocument:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX desteği için 'python-docx' kurulmalı: pip install -r requirements.txt") from exc

    document = Document(str(path))
    parts: list[str] = []
    paragraph_count = 0

    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if value:
            paragraph_count += 1
            parts.append(value)

    # Tablo içerikleri de sunum girdisine dahil edilir.
    for table_index, table in enumerate(document.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"[Tablo {table_index}]\n" + "\n".join(rows))

    text = _clean_text("\n\n".join(parts))
    if not text:
        raise ValueError("DOCX dosyasında kullanılabilir metin bulunamadı.")
    return ExtractedDocument(path, "docx", text, paragraph_count)


def extract_document(path: str | Path) -> ExtractedDocument:
    source = Path(path).expanduser().resolve()
    if not source.exists():
        raise FileNotFoundError(f"Dosya bulunamadı: {source}")
    if not source.is_file():
        raise ValueError(f"Girdi bir dosya olmalı: {source}")

    suffix = source.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(source)
    if suffix == ".docx":
        return _extract_docx(source)
    raise ValueError("Desteklenmeyen dosya türü. Şimdilik .pdf ve .docx destekleniyor.")
