from __future__ import annotations
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

BG = RGBColor(9, 12, 18)
CARD = RGBColor(17, 22, 31)
TEXT = RGBColor(240, 243, 248)
MUTED = RGBColor(163, 171, 184)
ACCENT = RGBColor(102, 241, 190)

def _add_text(slide, text, x, y, w, h, size, color=TEXT, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Arial"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box

def _background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG

def _footer(slide, index):
    _add_text(slide, f"{index:02d}", 12.2, 7.05, .6, .25, 8, MUTED)

def export_pptx(data, output_path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _background(slide)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.65), Inches(.75), Inches(.08), Inches(5.8))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
    _add_text(slide, data["title"], 1.05, 1.55, 10.8, 1.6, 30, TEXT, True)
    _add_text(slide, data.get("subtitle",""), 1.08, 3.25, 9.8, .8, 16, MUTED)
    _add_text(slide, "AI PRESENTATION AUTOMATION", 1.08, 5.85, 5.5, .3, 9, ACCENT, True)

    for i, s in enumerate(data["slides"], start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _background(slide)

        # Accent label
        _add_text(slide, f"SLIDE {i:02d}", .75, .55, 2, .3, 8, ACCENT, True)
        _add_text(slide, s["title"], .75, 1.0, 11.6, .9, 25, TEXT, True)
        _add_text(slide, s["key_message"], .78, 1.95, 10.8, .65, 13, MUTED)

        # Content card
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.75), Inches(2.85), Inches(7.65), Inches(3.55))
        shape.fill.solid(); shape.fill.fore_color.rgb = CARD
        shape.line.color.rgb = RGBColor(38, 45, 57)

        y = 3.18
        for bullet in s.get("content", [])[:5]:
            _add_text(slide, "•", 1.08, y, .3, .35, 14, ACCENT, True)
            _add_text(slide, bullet, 1.42, y-.02, 6.45, .55, 14, TEXT)
            y += .58

        # Visual direction card
        v = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.7), Inches(2.85), Inches(3.85), Inches(3.55))
        v.fill.solid(); v.fill.fore_color.rgb = CARD
        v.line.color.rgb = RGBColor(38, 45, 57)
        _add_text(slide, "VISUAL DIRECTION", 9.05, 3.18, 2.7, .3, 8, ACCENT, True)
        _add_text(slide, s.get("visual",""), 9.05, 3.7, 2.95, 1.65, 14, TEXT)
        _footer(slide, i)

        notes = s.get("speaker_notes")
        if notes:
            try:
                notes_slide = slide.notes_slide
                notes_slide.notes_text_frame.text = notes
            except Exception:
                pass

    prs.save(output_path)
    return output_path
